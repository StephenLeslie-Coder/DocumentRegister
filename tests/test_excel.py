import tempfile
import unittest
from datetime import date
from pathlib import Path

from src.excel import ExcelRegister
from src.models import DocumentResult, Status


class ExcelTests(unittest.TestCase):
    def test_appends_once_and_rejects_review(self):
        with tempfile.TemporaryDirectory() as directory:
            pdf = Path(directory) / "memo.pdf"
            pdf.write_bytes(b"first")
            register = ExcelRegister(Path(directory) / "register.xlsx")
            result = DocumentResult("memo.pdf", "Kamaya Thompson", "Hansel Ramdhan",
                                    "Procurement request", Status.SUCCESS,
                                    origin="Public Procurement Branch", receipt_date=date(2026, 10, 1))
            self.assertTrue(register.append(pdf, result))
            self.assertFalse(register.append(pdf, result))
            pdf.write_bytes(b"second")
            self.assertTrue(register.append(pdf, result))
            with self.assertRaises(ValueError):
                register.append(pdf, DocumentResult("memo.pdf"))
            from openpyxl import load_workbook
            book = load_workbook(register.path)
            try:
                self.assertEqual(book.active.cell(2, 2).value, "Public Procurement Branch")
                self.assertEqual(book.active.cell(2, 6).value.date(), date(2026, 10, 1))
                self.assertEqual(book.active.cell(2, 6).number_format, "dd-mmm-yyyy")
                self.assertTrue(book.active.column_dimensions["I"].hidden)
            finally:
                book.close()

    def test_migrates_legacy_register_without_losing_rows_or_keys(self):
        from openpyxl import Workbook, load_workbook
        from src.excel import LEGACY_HEADERS, HEADERS, document_key
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            pdf = root / "memo.pdf"
            pdf.write_bytes(b"old")
            register = ExcelRegister(root / "register.xlsx")
            book = Workbook()
            book.active.append(LEGACY_HEADERS)
            book.active.append(["memo.pdf", "Old To", "Old From", "Old Subject", None,
                                "SUCCESS", document_key(pdf)])
            book.active.column_dimensions["G"].hidden = True
            book.save(register.path)
            book.close()
            register.validate()
            self.assertTrue(register.contains(pdf))
            self.assertTrue(register.migrate())
            self.assertFalse(register.migrate())
            self.assertTrue(register.contains(pdf))
            book = load_workbook(register.path)
            try:
                sheet = book.active
                self.assertEqual([sheet.cell(1, c).value for c in range(1, 10)], HEADERS)
                self.assertEqual([sheet.cell(2, c).value for c in (1, 3, 4, 5, 8, 9)],
                                 ["memo.pdf", "Old To", "Old From", "Old Subject", "SUCCESS", document_key(pdf)])
                self.assertEqual(sheet.max_row, 2)
                self.assertTrue(sheet.column_dimensions["I"].hidden)
                self.assertFalse(sheet.column_dimensions["G"].hidden)
            finally:
                book.close()


if __name__ == "__main__":
    unittest.main()

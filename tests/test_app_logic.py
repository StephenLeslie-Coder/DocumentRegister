import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from src.app_logic import (AppSettings, DisplayStatus, SettingsStore,
                           approve_and_save, discover_pdfs, process_one)
from src.excel import ExcelRegister
from src.models import DocumentResult, Status


class AppLogicTests(unittest.TestCase):
    def test_folder_discovery_is_direct_and_pdf_only(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            (folder / "b.PDF").write_bytes(b"pdf")
            (folder / "a.pdf").write_bytes(b"pdf")
            (folder / "notes.txt").write_text("ignore")
            (folder / "nested").mkdir()
            (folder / "nested" / "hidden.pdf").write_bytes(b"pdf")
            self.assertEqual([p.name for p in discover_pdfs(folder)], ["a.pdf", "b.PDF"])

    def test_settings_persist_paths_only(self):
        with tempfile.TemporaryDirectory() as directory:
            store = SettingsStore(Path(directory) / "settings.json")
            expected = AppSettings("C:/Inbox", "C:/Register.xlsx")
            store.save(expected)
            self.assertEqual(store.load(), expected)
            self.assertEqual(set(json.loads(store.path.read_text())), {"document_folder", "register_path"})

    def test_register_creation_review_approval_and_duplicate(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            pdf = root / "memo.pdf"
            pdf.write_bytes(b"sample")
            register = ExcelRegister(root / "register.xlsx")
            register.create()
            register.validate()
            with self.assertRaises(FileExistsError):
                register.create()
            with self.assertRaises(ValueError):
                approve_and_save(pdf, register, "Name", "  ", "Subject")
            self.assertFalse(register.contains(pdf))
            self.assertTrue(approve_and_save(pdf, register, " Corrected To ", "Corrected From", "Corrected\n  Subject"))
            self.assertTrue(register.contains(pdf))
            self.assertFalse(approve_and_save(pdf, register, "Corrected To", "Corrected From", "Corrected Subject"))
            book = load_workbook(register.path)
            try:
                self.assertEqual(book.active.max_row, 2)
                self.assertEqual([book.active.cell(2, i).value for i in (2, 3, 4)],
                                 ["Corrected To", "Corrected From", "Corrected Subject"])
            finally:
                book.close()

    def test_processing_writes_success_skips_review_and_detects_duplicate(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            pdf = root / "memo.pdf"
            pdf.write_bytes(b"sample")
            register = ExcelRegister(root / "register.xlsx")
            register.create()
            review = DocumentResult(pdf.name, "To Name", "From Name", "Subject", Status.NEEDS_REVIEW)
            success = DocumentResult(pdf.name, "To Name", "From Name", "Subject", Status.SUCCESS)
            with patch("src.app_logic.process_document", return_value=review):
                self.assertEqual(process_one(pdf, register, None).status, DisplayStatus.NEEDS_REVIEW)
            self.assertFalse(register.contains(pdf))
            with patch("src.app_logic.process_document", return_value=success) as extractor:
                self.assertEqual(process_one(pdf, register, None).status, DisplayStatus.PROCESSED)
                self.assertEqual(process_one(pdf, register, None).status, DisplayStatus.ALREADY_PROCESSED)
                extractor.assert_called_once()

    def test_missing_register_does_not_silently_create_one(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            pdf = root / "memo.pdf"
            pdf.write_bytes(b"sample")
            register = ExcelRegister(root / "missing.xlsx")
            outcome = process_one(pdf, register, None)
            self.assertEqual(outcome.status, DisplayStatus.FAILED)
            self.assertFalse(register.path.exists())


if __name__ == "__main__":
    unittest.main()

"""Append reviewed results while preserving legacy workbook rows and document keys."""

import hashlib
import logging
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from .models import DocumentResult, Status


LEGACY_HEADERS = ["File Name", "To", "From", "Subject", "Processed Date", "Status", "Document Key"]
HEADERS = ["File Name", "Origin", "To", "From", "Subject", "Receipt Date",
           "Processed Date", "Status", "Document Key"]
log = logging.getLogger(__name__)


def document_key(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return f"{path.name.casefold()}:{digest.hexdigest()}"


def _safe(value: str) -> str:
    # Prevent spreadsheet formula interpretation of text from an untrusted PDF.
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def _save_existing_atomically(book, path: Path) -> None:
    # Keep the original register intact if saving or replacing a migration fails.
    with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as target:
        temporary = Path(target.name)
    try:
        book.save(temporary)
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


class ExcelRegister:
    def __init__(self, path: Path):
        self.path = path

    @staticmethod
    def _new_book():
        book = Workbook()
        sheet = book.active
        sheet.title = "Register"
        sheet.append(HEADERS)
        sheet.column_dimensions["I"].hidden = True
        for column, width in {"A": 30, "B": 42, "C": 32, "D": 32, "E": 64,
                              "F": 19, "G": 22, "H": 18}.items():
            sheet.column_dimensions[column].width = width
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = "A1:H1"
        return book

    @staticmethod
    def _check_headers(sheet) -> bool:
        if [sheet.cell(1, c).value for c in range(1, 10)] == HEADERS:
            return False
        if [sheet.cell(1, c).value for c in range(1, 8)] == LEGACY_HEADERS:
            return True
        raise ValueError("Register has an unexpected header; refusing to modify it")

    @staticmethod
    def _migrate_sheet(sheet) -> None:
        sheet.insert_cols(2)
        sheet.insert_cols(6)
        sheet.cell(1, 2, "Origin")
        sheet.cell(1, 6, "Receipt Date")
        for column, width in {"A": 30, "B": 42, "C": 32, "D": 32, "E": 64,
                              "F": 19, "G": 22, "H": 18}.items():
            sheet.column_dimensions[column].width = width
        sheet.column_dimensions["G"].hidden = False
        sheet.column_dimensions["I"].hidden = True
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.auto_filter.ref = "A1:H1"

    def migrate(self) -> bool:
        book = load_workbook(self.path)
        try:
            if not self._check_headers(book.active):
                return False
            self._migrate_sheet(book.active)
            _save_existing_atomically(book, self.path)
            return True
        finally:
            book.close()

    def create(self) -> None:
        if self.path.exists():
            raise FileExistsError(self.path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        book = self._new_book()
        try:
            book.save(self.path)
        finally:
            book.close()

    def validate(self) -> None:
        book = load_workbook(self.path, read_only=True)
        try:
            self._check_headers(book.active)
        finally:
            book.close()

    def contains(self, pdf_path: Path) -> bool:
        if not self.path.exists():
            return False
        key = document_key(pdf_path)
        book = load_workbook(self.path, read_only=True)
        try:
            sheet = book.active
            legacy = self._check_headers(sheet)
            column = 7 if legacy else 9
            return any(row[0] == key for row in sheet.iter_rows(min_row=2, min_col=column,
                                                                 max_col=column, values_only=True))
        finally:
            book.close()

    def append(self, pdf_path: Path, result: DocumentResult) -> bool:
        if result.status != Status.SUCCESS:
            raise ValueError("Only reviewed SUCCESS results can be added to the register")
        if not all((result.origin, result.to, result.from_, result.subject, result.receipt_date)):
            raise ValueError("Origin, To, From, Subject, and Receipt Date are required")
        key = document_key(pdf_path)
        if self.path.exists():
            book = load_workbook(self.path)
        else:
            book = self._new_book()
        try:
            sheet = book.active
            legacy = self._check_headers(sheet)
            if legacy:
                self._migrate_sheet(sheet)
            if any(sheet.cell(row, 9).value == key for row in range(2, sheet.max_row + 1)):
                if legacy:
                    _save_existing_atomically(book, self.path)
                return False
            sheet.append([_safe(result.file_name), _safe(result.origin), _safe(result.to),
                          _safe(result.from_), _safe(result.subject), result.receipt_date,
                          datetime.now(), result.status.value, key])
            sheet.cell(sheet.max_row, 6).number_format = "dd-mmm-yyyy"
            self.path.parent.mkdir(parents=True, exist_ok=True)
            if self.path.exists():
                _save_existing_atomically(book, self.path)
            else:
                book.save(self.path)
            log.info("Excel row created for %s", pdf_path.name)
            return True
        finally:
            book.close()
